import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


PALETTE = {
    "ink": "1C1A17",
    "muted": "746A5C",
    "paper": "F4EFE6",
    "panel": "FBF8F2",
    "border": "D9D0C1",
    "accent": "C86D44",
    "blue": "355E8C",
    "white": "FFFFFF",
}

THIN = Side(style="thin", color=PALETTE["border"])


def load_payload(path_str):
    with open(path_str, "r", encoding="utf-8") as handle:
        return json.load(handle)


def apply_sheet_defaults(ws):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = PALETTE["accent"]


def style_cell(cell, *, font=None, fill=None, alignment=None, border=None, number_format=None):
    if font is not None:
        cell.font = font
    if fill is not None:
        cell.fill = fill
    if alignment is not None:
        cell.alignment = alignment
    if border is not None:
        cell.border = border
    if number_format is not None:
        cell.number_format = number_format


def title_block(ws, title, subtitle, span=8):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span)
    ws["A1"] = title
    style_cell(
        ws["A1"],
        font=Font(bold=True, size=16, color=PALETTE["ink"]),
        alignment=Alignment(horizontal="left"),
    )
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=span)
    ws["A2"] = subtitle
    style_cell(
        ws["A2"],
        font=Font(size=10, color=PALETTE["muted"]),
        alignment=Alignment(horizontal="left"),
    )


def section_header(ws, row, label, span=8):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    cell = ws.cell(row=row, column=1, value=label)
    style_cell(
        cell,
        font=Font(bold=True, size=11, color=PALETTE["white"]),
        fill=PatternFill("solid", fgColor=PALETTE["blue"]),
        alignment=Alignment(horizontal="left"),
        border=Border(left=THIN, right=THIN, top=THIN, bottom=THIN),
    )


def table_header(ws, row, headers):
    fill = PatternFill("solid", fgColor=PALETTE["accent"])
    font = Font(bold=True, color=PALETTE["white"])
    alignment = Alignment(horizontal="center", vertical="center")
    border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=header)
        style_cell(cell, font=font, fill=fill, alignment=alignment, border=border)


def body_cell(ws, row, col, value, *, number_format=None, align="left"):
    cell = ws.cell(row=row, column=col, value=value)
    style_cell(
        cell,
        font=Font(size=10, color=PALETTE["ink"]),
        fill=PatternFill("solid", fgColor=PALETTE["panel"]),
        alignment=Alignment(horizontal=align, vertical="top", wrap_text=True),
        border=Border(left=THIN, right=THIN, top=THIN, bottom=THIN),
        number_format=number_format,
    )
    return cell


def autosize(ws, widths):
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def write_cover(ws, payload):
    apply_sheet_defaults(ws)
    meta = payload.get("meta", {})
    summary = payload.get("summary", {})
    title_block(
        ws,
        f'{meta.get("brand", "Keystone AI")} | Planning Estimate',
        f'{meta.get("projectName", "Project")} | Generated {str(meta.get("generatedAt", ""))[:19].replace("T", " ")}',
        span=6,
    )

    row = 4
    section_header(ws, row, "Project Overview", span=6)
    row += 1
    overview_rows = [
        ("Project Name", meta.get("projectName", "")),
        ("Estimate Type", str(meta.get("estimateType", "concept")).title()),
        ("Budget Tier", meta.get("budgetTier", "")),
        ("Rate Family", meta.get("rateFamily", "")),
        ("Front Facing", meta.get("frontFacing", "")),
        ("Lot Context", meta.get("lotContext", "")),
        ("Conditioned Area (sqft)", summary.get("conditionedAreaSqFt", 0)),
        ("Garage Area (sqft)", summary.get("garageAreaSqFt", 0)),
    ]
    for label, value in overview_rows:
        body_cell(ws, row, 1, label)
        if isinstance(value, (int, float)):
            body_cell(ws, row, 2, value, number_format='#,##0.0', align="right")
        else:
            body_cell(ws, row, 2, value)
        row += 1

    row += 1
    section_header(ws, row, "Planning Cost Range", span=6)
    row += 1
    table_header(ws, row, ["Low", "Target", "High"])
    row += 1
    totals = payload.get("totals", {})
    body_cell(ws, row, 1, totals.get("low", 0), number_format='$#,##0', align="right")
    body_cell(ws, row, 2, totals.get("target", 0), number_format='$#,##0', align="right")
    body_cell(ws, row, 3, totals.get("high", 0), number_format='$#,##0', align="right")

    row += 2
    section_header(ws, row, "Survey Highlights", span=6)
    row += 1
    table_header(ws, row, ["Attribute", "Value"])
    row += 1
    for item in payload.get("surveyHighlights", []):
        body_cell(ws, row, 1, item.get("label", ""))
        body_cell(ws, row, 2, item.get("value", ""))
        row += 1

    autosize(ws, {"A": 30, "B": 34, "C": 16, "D": 16, "E": 16, "F": 16})


def write_summary(ws, payload):
    apply_sheet_defaults(ws)
    title_block(ws, "Summary", "Category totals and range composition", span=6)
    table_header(ws, 3, ["Category", "Low", "Target", "High", "Target Share"])
    ws.freeze_panes = "A4"

    row = 4
    total_target = float(payload.get("totals", {}).get("target", 0) or 0)
    for category in payload.get("categoryTotals", []):
        target_value = float(category.get("target", 0) or 0)
        share = (target_value / total_target) if total_target > 0 else 0
        body_cell(ws, row, 1, category.get("label", ""))
        body_cell(ws, row, 2, category.get("low", 0), number_format='$#,##0', align="right")
        body_cell(ws, row, 3, target_value, number_format='$#,##0', align="right")
        body_cell(ws, row, 4, category.get("high", 0), number_format='$#,##0', align="right")
        body_cell(ws, row, 5, share, number_format='0.0%', align="right")
        row += 1

    row += 1
    table_header(ws, row, ["Total Low", "Total Target", "Total High"])
    row += 1
    totals = payload.get("totals", {})
    body_cell(ws, row, 1, totals.get("low", 0), number_format='$#,##0', align="right")
    body_cell(ws, row, 2, totals.get("target", 0), number_format='$#,##0', align="right")
    body_cell(ws, row, 3, totals.get("high", 0), number_format='$#,##0', align="right")
    autosize(ws, {"A": 30, "B": 16, "C": 16, "D": 16, "E": 16, "F": 16})


def write_detail(ws, payload):
    apply_sheet_defaults(ws)
    title_block(ws, "Detail", "Line-item material and labor breakdown", span=8)
    table_header(ws, 3, ["Category", "Line Item", "Qty", "Unit", "Material Cost", "Labor Cost", "Total"])
    ws.freeze_panes = "A4"

    row = 4
    for item in payload.get("detailLineItems", []):
        body_cell(ws, row, 1, item.get("category", ""))
        body_cell(ws, row, 2, item.get("lineItem", ""))
        body_cell(ws, row, 3, item.get("quantity", 0), number_format='#,##0.0', align="right")
        body_cell(ws, row, 4, item.get("unit", ""), align="center")
        body_cell(ws, row, 5, item.get("materialCost", 0), number_format='$#,##0', align="right")
        body_cell(ws, row, 6, item.get("laborCost", 0), number_format='$#,##0', align="right")
        body_cell(ws, row, 7, item.get("total", 0), number_format='$#,##0', align="right")
        row += 1

    autosize(ws, {"A": 24, "B": 36, "C": 12, "D": 10, "E": 16, "F": 16, "G": 16, "H": 16})


def write_takeoff(ws, payload):
    apply_sheet_defaults(ws)
    title_block(ws, "Takeoff", "Room schedule and quantity takeoff", span=10)
    table_header(
        ws,
        3,
        ["Level", "Room", "Requested Label", "Canonical Type", "Conditioned", "W (ft)", "H (ft)", "Area (sqft)", "X", "Y", "Ceiling (ft)", "Finish Class"],
    )
    ws.freeze_panes = "A4"

    row = 4
    for item in payload.get("roomSchedule", []):
        body_cell(ws, row, 1, item.get("level", ""), align="center")
        body_cell(ws, row, 2, item.get("label", ""))
        body_cell(ws, row, 3, item.get("requestedLabel", ""))
        body_cell(ws, row, 4, item.get("canonicalType", ""))
        body_cell(ws, row, 5, "Yes" if item.get("conditioned") else "No", align="center")
        body_cell(ws, row, 6, item.get("widthFt", 0), number_format='#,##0', align="right")
        body_cell(ws, row, 7, item.get("heightFt", 0), number_format='#,##0', align="right")
        body_cell(ws, row, 8, item.get("areaSqFt", 0), number_format='#,##0.0', align="right")
        body_cell(ws, row, 9, item.get("x", 0), number_format='#,##0', align="right")
        body_cell(ws, row, 10, item.get("y", 0), number_format='#,##0', align="right")
        body_cell(ws, row, 11, item.get("ceilingHeightFt", 0), number_format='#,##0.0', align="right")
        body_cell(ws, row, 12, item.get("finishClass", ""))
        row += 1

    row += 2
    section_header(ws, row, "Assemblies", span=10)
    row += 1
    table_header(ws, row, ["Item", "Unit", "Quantity", "Notes"])
    row += 1
    for item in payload.get("assemblies", []):
        body_cell(ws, row, 1, item.get("label", ""))
        body_cell(ws, row, 2, item.get("unit", ""), align="center")
        body_cell(ws, row, 3, item.get("quantity", 0), number_format='#,##0.0', align="right")
        body_cell(ws, row, 4, item.get("notes", ""))
        row += 1

    autosize(ws, {"A": 12, "B": 28, "C": 24, "D": 20, "E": 12, "F": 10, "G": 10, "H": 14, "I": 8, "J": 8, "K": 12, "L": 18})


def write_assumptions(ws, payload):
    apply_sheet_defaults(ws)
    assumptions = payload.get("assumptions", {})
    title_block(ws, "Assumptions", "Model inputs, material selections, and notes", span=8)

    row = 4
    section_header(ws, row, "Assumption Inputs", span=8)
    row += 1
    key_rows = [
        ("Currency", assumptions.get("currency", "USD")),
        ("Budget Tier", assumptions.get("budgetTier", "")),
        ("Rate Family", assumptions.get("rateFamily", "")),
        ("Regional Multiplier", assumptions.get("regionalMultiplier", 1)),
        ("Roof Kind", assumptions.get("roofKind", "")),
        ("Roof Area Multiplier", assumptions.get("roofAreaMultiplier", 0)),
        ("Ceiling Height Type", assumptions.get("ceilingHeightType", "")),
    ]
    for label, value in key_rows:
        body_cell(ws, row, 1, label)
        if isinstance(value, (int, float)):
            body_cell(ws, row, 2, value, number_format='#,##0.00', align="right")
        else:
            body_cell(ws, row, 2, value)
        row += 1

    row += 1
    section_header(ws, row, "Material Selections", span=8)
    row += 1
    table_header(ws, row, ["Selection", "Value"])
    row += 1
    for item in payload.get("materialSelections", []):
        body_cell(ws, row, 1, item.get("label", ""))
        body_cell(ws, row, 2, item.get("value", ""))
        row += 1

    row += 1
    section_header(ws, row, "Data Sources", span=8)
    row += 1
    for note in assumptions.get("dataSources", []):
        body_cell(ws, row, 1, note)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        row += 1

    row += 1
    section_header(ws, row, "Notes", span=8)
    row += 1
    for note in assumptions.get("notes", []):
        body_cell(ws, row, 1, note)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        row += 1

    autosize(ws, {"A": 34, "B": 32, "C": 16, "D": 16, "E": 16, "F": 16, "G": 16, "H": 16})


def write_alternates(ws, payload):
    apply_sheet_defaults(ws)
    title_block(ws, "Alternates", "Value-engineering options for cost reduction", span=8)
    table_header(ws, 3, ["Scope", "Current Option", "Alternate Option", "Quantity", "Unit", "Est. Savings", "Notes"])
    ws.freeze_panes = "A4"

    row = 4
    alternates = payload.get("alternates", [])
    if not alternates:
        body_cell(ws, row, 1, "No value-engineering alternates available for current selections.")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        autosize(ws, {"A": 30, "B": 30, "C": 30, "D": 12, "E": 10, "F": 14, "G": 36, "H": 16})
        return

    total_savings = 0
    for item in alternates:
        body_cell(ws, row, 1, item.get("scope", ""))
        body_cell(ws, row, 2, item.get("currentOption", ""))
        body_cell(ws, row, 3, item.get("alternateOption", ""))
        body_cell(ws, row, 4, item.get("quantity", 0), number_format='#,##0.0', align="right")
        body_cell(ws, row, 5, item.get("unit", ""), align="center")
        savings = item.get("estimatedSavings", 0) or 0
        body_cell(ws, row, 6, savings, number_format='$#,##0', align="right")
        body_cell(ws, row, 7, item.get("notes", ""))
        total_savings += savings
        row += 1

    row += 1
    table_header(ws, row, ["Total Potential Savings"])
    row += 1
    body_cell(ws, row, 1, total_savings, number_format='$#,##0', align="right")
    autosize(ws, {"A": 30, "B": 30, "C": 30, "D": 12, "E": 10, "F": 14, "G": 36, "H": 16})


def build_workbook(payload):
    wb = Workbook()

    cover_ws = wb.active
    cover_ws.title = "Cover"
    write_cover(cover_ws, payload)

    summary_ws = wb.create_sheet("Summary")
    write_summary(summary_ws, payload)

    detail_ws = wb.create_sheet("Detail")
    write_detail(detail_ws, payload)

    takeoff_ws = wb.create_sheet("Takeoff")
    write_takeoff(takeoff_ws, payload)

    assumptions_ws = wb.create_sheet("Assumptions")
    write_assumptions(assumptions_ws, payload)

    alternates_ws = wb.create_sheet("Alternates")
    write_alternates(alternates_ws, payload)

    return wb


def main():
    if len(sys.argv) != 3:
        raise SystemExit("Usage: build_estimate_xlsx.py <input.json> <output.xlsx>")

    input_path = Path(sys.argv[1])
    output_path = Path(sys.argv[2])
    payload = load_payload(str(input_path))
    workbook = build_workbook(payload)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(str(output_path))


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        sys.stderr.write(str(exc))
        sys.exit(1)
